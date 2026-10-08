"""Read a returned native-speaker review sheet (Word) and turn it into project data.

  python scripts/ingest_review.py path/to/AHUIKE_Igbo_Review_Sheet_filled.docx --lang ig

Writes:
  data/advice_reviewed.json - Task 1: reviewed advice per language. Keys the reviewer marked wrong without a
                              correction (or that are --hold) are stored as "" so the app falls back to
                              English instead of showing unreviewed machine translation
  data/native_items.jsonl   - Task 3: cases written by the native speaker (variant "native"), benchmarked separately
  data/verified_<lang>.json - Task 4: SAME / SMALL / WRONG verdict per AI-translated case
  data/review_<lang>.json   - everything read, including glossary answers (Task 2), for the record
Re-running replaces this language's entries, so a corrected sheet can simply be ingested again.
"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")
sys.path.insert(0, str(ROOT))

from ahuike.prompts import ADVICE_EN  # noqa: E402


def cell_text(cell) -> str:
    return "\n".join(p.text for p in cell.paragraphs).strip()


def tables_by_header(doc) -> dict[str, list[list[str]]]:
    """Map a recognisable header phrase -> table rows (excluding the header row)."""
    found = {}
    for t in doc.tables:
        rows = [[cell_text(c) for c in r.cells] for r in t.rows]
        if not rows:
            continue
        head = " | ".join(rows[0]).lower()
        if "english (what it must say)" in head:
            found["advice"] = rows[1:]
        elif "igbo word we use" in head or "word we use" in head:
            found["glossary"] = rows[1:]
        elif "facts to include" in head:
            found["native"] = rows[1:]
        elif "meaning" in head and "english" in head:
            found["verify"] = rows[1:]
    return found


def norm_verdict(s: str, options: dict[str, str]) -> str | None:
    s = s.strip().upper()
    for key, val in options.items():
        if s.startswith(key):
            return val
    return None


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("docx")
    ap.add_argument("--lang", default="ig", choices=("ha", "yo", "ig"))
    ap.add_argument("--data", default=str(ROOT / "data"))
    ap.add_argument("--hold", nargs="*", default=[],
                    help="advice keys (e.g. child|EMERGENCY_REFER_NOW) to withhold pending confirmation")
    args = ap.parse_args()

    import docx  # python-docx

    data = Path(args.data)
    tabs = tables_by_header(docx.Document(args.docx))
    report: dict = {"source": Path(args.docx).name, "lang": args.lang}

    # ---- Task 1: advice. Row order matches ADVICE_EN (child, pregnant, postpartum x E, C, H).
    advice_path = data / "advice_reviewed.json"
    table = json.loads(advice_path.read_text(encoding="utf-8")) if advice_path.exists() else {}
    keys = [f"{g}|{lv}" for (g, lv) in ADVICE_EN]
    adv_rows = tabs.get("advice", [])
    advice_report = []
    for key, row in zip(keys, adv_rows):
        machine, verdict_raw, corrected = row[3], row[4], row[5]
        verdict = norm_verdict(verdict_raw, {"C": "correct", "S": "small_fix", "W": "wrong"})
        # Reviewers sometimes append an explanation ("... Review: iri means eat"); keep it as a note.
        note = ""
        for marker in ("Review:", "review:", "Note:", "note:"):
            if marker in corrected:
                corrected, note = (x.strip() for x in corrected.split(marker, 1))
                break
        # The reviewer's own wording wins; a "correct" verdict keeps N-ATLAS's text.
        final = corrected or (machine.split(" Check:")[0].strip() if verdict == "correct" else None)
        held = key in args.hold
        advice_report.append({"key": key, "verdict": verdict, "corrected": bool(corrected), "note": note,
                              "held": held, "final": None if held else final})
        # "" = do not show this language for this message (English fallback) until resolved.
        table.setdefault(args.lang, {})[key] = "" if held or not final else final
    if adv_rows:
        advice_path.write_text(json.dumps(table, ensure_ascii=False, indent=2), encoding="utf-8")
    report["advice"] = advice_report
    unresolved = [a["key"] for a in advice_report if not a["final"]]  # includes held keys

    # ---- Task 2: glossary answers (applied to the translator glossary by hand)
    report["glossary"] = [{"english": r[0], "ours": r[1], "check": r[2], "mothers_use": r[3]}
                          for r in tabs.get("glossary", []) if any(r[2:])]

    # ---- Task 3: native-written cases
    native_path = data / "native_items.jsonl"
    keep = []
    if native_path.exists():
        keep = [json.loads(l) for l in native_path.read_text(encoding="utf-8").splitlines() if l.strip()]
        keep = [r for r in keep if r["lang"] != args.lang]
    native = [{"case_id": r[1], "lang": args.lang, "variant": "native", "text": r[4]}
              for r in tabs.get("native", []) if len(r) >= 5 and r[1] and r[4]]
    with open(native_path, "w", encoding="utf-8") as fh:
        for r in keep + native:
            fh.write(json.dumps(r, ensure_ascii=False) + "\n")
    report["native_cases"] = len(native)

    # ---- Task 4: verdicts on AI translations
    verdicts = {}
    for r in tabs.get("verify", []):
        if len(r) >= 6 and r[1]:
            v = norm_verdict(r[4], {"SAME": "same", "SMALL": "small", "WRONG": "wrong"})
            if v:
                verdicts[r[1]] = {"verdict": v, "note": r[5]}
    (data / f"verified_{args.lang}.json").write_text(json.dumps(verdicts, ensure_ascii=False, indent=1), encoding="utf-8")
    report["verified"] = {v: sum(x["verdict"] == v for x in verdicts.values()) for v in ("same", "small", "wrong")}

    (data / f"review_{args.lang}.json").write_text(json.dumps(report, ensure_ascii=False, indent=1), encoding="utf-8")
    print(json.dumps({k: report[k] for k in ("native_cases", "verified")}, indent=1))
    print(f"advice: {sum(a['final'] is not None for a in advice_report)}/9 resolved"
          + (f"; STILL NEEDED: {unresolved}" if unresolved else ""))
    if report["glossary"]:
        print("glossary feedback (update GLOSSARY in scripts/translate_natlas.py):")
        for g in report["glossary"]:
            print("  ", g)


if __name__ == "__main__":
    main()
